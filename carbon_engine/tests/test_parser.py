"""
Unit tests for the Tefnut data ingestion parser.
Uses in-memory Excel/CSV construction — no file I/O needed.
"""

from __future__ import annotations

import io

import openpyxl
import pandas as pd
import pytest

from carbon_engine.parser import (
    TemplateParseError,
    parse_csv,
    parse_excel,
    parse_file,
    PI_PLANT_NAME,
    PI_REPORTING_YEAR,
    PI_CALCINATION_METHOD,
    PI_CEMENT_PRODUCTION,
    PI_B1_CLINKER,
    PI_ELEC_MWH,
    PI_ELEC_EF,
    PI_WATER_WITHDRAWAL,
    PI_WATER_DISCHARGE,
    PI_WATER_RECYCLED,
    FUEL_TYPE_COL,
    FUEL_CONSUMPTION_COL,
    FUEL_LHV_COL,
    FUEL_EF_COL,
    FUEL_BIO_COL,
    TRANSPORT_DESC_COL,
    TRANSPORT_TKM_COL,
    TRANSPORT_EF_COL,
)
from carbon_engine.carbon_engine import CalcinationMethod


# ---------------------------------------------------------------------------
# Helpers to build in-memory Excel workbooks
# ---------------------------------------------------------------------------
def _make_workbook(
    plant_row: dict,
    kiln_fuels: list[dict] | None = None,
    non_kiln_fuels: list[dict] | None = None,
    transport: list[dict] | None = None,
) -> bytes:
    """Build a minimal valid Tefnut Excel workbook in memory."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # plant_info sheet — row 1 = column names, row 2 = data
    ws_pi = wb.create_sheet("plant_info")
    headers = list(plant_row.keys())
    for col_idx, h in enumerate(headers, start=1):
        ws_pi.cell(row=1, column=col_idx, value=h)
    for col_idx, v in enumerate(plant_row.values(), start=1):
        ws_pi.cell(row=2, column=col_idx, value=v)

    def _write_sheet(name: str, rows: list[dict]) -> None:
        ws = wb.create_sheet(name)
        if not rows:
            return
        hdrs = list(rows[0].keys())
        for col_idx, h in enumerate(hdrs, start=1):
            ws.cell(row=1, column=col_idx, value=h)
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, v in enumerate(row.values(), start=1):
                ws.cell(row=row_idx, column=col_idx, value=v)

    if kiln_fuels is not None:
        _write_sheet("kiln_fuels", kiln_fuels)
    if non_kiln_fuels is not None:
        _write_sheet("non_kiln_fuels", non_kiln_fuels)
    if transport is not None:
        _write_sheet("transport", transport)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _base_plant_row(**overrides) -> dict:
    base = {
        PI_PLANT_NAME: "Test Plant",
        PI_REPORTING_YEAR: 2026,
        PI_CALCINATION_METHOD: "B1",
        PI_CEMENT_PRODUCTION: 1_200_000,
        PI_B1_CLINKER: 1_000_000,
    }
    base.update(overrides)
    return base


def _base_fuel_row(**overrides) -> dict:
    base = {
        FUEL_TYPE_COL: "coal_anthracite",
        FUEL_CONSUMPTION_COL: 10_000,
        FUEL_LHV_COL: 26.0,
        FUEL_EF_COL: None,
        FUEL_BIO_COL: 0.0,
    }
    base.update(overrides)
    return base


# ---------------------------------------------------------------------------
# Excel parsing — happy path
# ---------------------------------------------------------------------------
class TestParseExcelHappyPath:
    def test_minimal_b1(self):
        wb_bytes = _make_workbook(_base_plant_row())
        plant = parse_excel(wb_bytes)
        assert plant.plant_name == "Test Plant"
        assert plant.reporting_year == 2026
        assert plant.calcination_method == CalcinationMethod.B1_SIMPLE_OUTPUT
        assert plant.calcination_b1 is not None
        assert plant.calcination_b1.clinker_production_t_yr == 1_000_000

    def test_with_kiln_fuels(self):
        wb_bytes = _make_workbook(
            _base_plant_row(),
            kiln_fuels=[_base_fuel_row(), _base_fuel_row(fuel_type="natural_gas", lhv_gj_per_t=48.0)],
        )
        plant = parse_excel(wb_bytes)
        assert len(plant.kiln_fuels) == 2
        assert plant.kiln_fuels[0].fuel_type.value == "coal_anthracite"
        assert plant.kiln_fuels[1].fuel_type.value == "natural_gas"

    def test_default_ef_applied(self):
        """When ef_kg_co2_per_gj is blank, GCCA default should be applied."""
        wb_bytes = _make_workbook(
            _base_plant_row(),
            kiln_fuels=[_base_fuel_row(ef_kg_co2_per_gj=None)],
        )
        plant = parse_excel(wb_bytes)
        assert plant.kiln_fuels[0].ef_kg_co2_per_gj == 96.0  # GCCA default for coal

    def test_custom_ef_preserved(self):
        wb_bytes = _make_workbook(
            _base_plant_row(),
            kiln_fuels=[_base_fuel_row(ef_kg_co2_per_gj=90.0)],
        )
        plant = parse_excel(wb_bytes)
        assert plant.kiln_fuels[0].ef_kg_co2_per_gj == 90.0

    def test_electricity_parsed(self):
        row = _base_plant_row(**{PI_ELEC_MWH: 50_000, PI_ELEC_EF: 0.5})
        wb_bytes = _make_workbook(row)
        plant = parse_excel(wb_bytes)
        assert plant.electricity is not None
        assert plant.electricity.purchased_electricity_mwh_yr == 50_000
        assert plant.electricity.grid_ef_kg_co2_per_mwh == 0.5

    def test_water_parsed(self):
        row = _base_plant_row(**{
            PI_WATER_WITHDRAWAL: 1_200_000,
            PI_WATER_DISCHARGE: 200_000,
            PI_WATER_RECYCLED: 50_000,
        })
        wb_bytes = _make_workbook(row)
        plant = parse_excel(wb_bytes)
        assert plant.water is not None
        assert plant.water.withdrawal_m3_yr == 1_200_000

    def test_transport_parsed(self):
        transport = [
            {TRANSPORT_DESC_COL: "Limestone", TRANSPORT_TKM_COL: 1_000_000, TRANSPORT_EF_COL: 0.062},
        ]
        wb_bytes = _make_workbook(_base_plant_row(), transport=transport)
        plant = parse_excel(wb_bytes)
        assert len(plant.transport_entries) == 1
        assert plant.transport_entries[0].freight_tonne_km == 1_000_000

    def test_transport_default_ef(self):
        """ef_kg_co2_per_tonne_km should default to 0.062 if not provided."""
        transport = [{TRANSPORT_DESC_COL: "Coal", TRANSPORT_TKM_COL: 500_000}]
        wb_bytes = _make_workbook(_base_plant_row(), transport=transport)
        plant = parse_excel(wb_bytes)
        assert plant.transport_entries[0].ef_kg_co2_per_tonne_km == 0.062

    def test_non_kiln_fuels_parsed(self):
        wb_bytes = _make_workbook(
            _base_plant_row(),
            non_kiln_fuels=[_base_fuel_row(fuel_type="diesel_oil", lhv_gj_per_t=42.7)],
        )
        plant = parse_excel(wb_bytes)
        assert len(plant.non_kiln_fuels) == 1
        assert plant.non_kiln_fuels[0].fuel_type.value == "diesel_oil"

    def test_a1_method(self):
        row = {
            PI_PLANT_NAME: "Plant A1",
            PI_REPORTING_YEAR: 2026,
            PI_CALCINATION_METHOD: "A1",
            PI_CEMENT_PRODUCTION: 1_200_000,
            "a1_kiln_feed_t_yr": 1_861_500,
            "a1_dust_return_correction_pct": 10.33,
            "a1_loi_raw_meal_pct": 35.0,
        }
        wb_bytes = _make_workbook(row)
        plant = parse_excel(wb_bytes)
        assert plant.calcination_method == CalcinationMethod.A1_SIMPLE_INPUT
        assert plant.calcination_a1 is not None
        assert plant.calcination_a1.loi_raw_meal_pct == 35.0


# ---------------------------------------------------------------------------
# Excel parsing — error cases
# ---------------------------------------------------------------------------
class TestParseExcelErrors:
    def test_missing_plant_info_sheet(self):
        wb = openpyxl.Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(TemplateParseError) as exc_info:
            parse_excel(buf.getvalue())
        assert any("plant_info" in e.message for e in exc_info.value.errors)

    def test_missing_required_plant_name(self):
        row = _base_plant_row()
        del row[PI_PLANT_NAME]
        wb_bytes = _make_workbook(row)
        with pytest.raises(TemplateParseError) as exc_info:
            parse_excel(wb_bytes)
        assert any(PI_PLANT_NAME in e.column for e in exc_info.value.errors)

    def test_missing_b1_clinker(self):
        row = _base_plant_row()
        del row[PI_B1_CLINKER]
        wb_bytes = _make_workbook(row)
        with pytest.raises(TemplateParseError):
            parse_excel(wb_bytes)

    def test_invalid_fuel_type(self):
        wb_bytes = _make_workbook(
            _base_plant_row(),
            kiln_fuels=[_base_fuel_row(fuel_type="rocket_fuel")],
        )
        with pytest.raises(TemplateParseError) as exc_info:
            parse_excel(wb_bytes)
        assert any("rocket_fuel" in e.message for e in exc_info.value.errors)

    def test_invalid_calcination_method(self):
        row = _base_plant_row(**{PI_CALCINATION_METHOD: "C3"})
        wb_bytes = _make_workbook(row)
        with pytest.raises(TemplateParseError) as exc_info:
            parse_excel(wb_bytes)
        assert any("C3" in e.message for e in exc_info.value.errors)

    def test_a1_missing_required_fields(self):
        row = {
            PI_PLANT_NAME: "Plant",
            PI_REPORTING_YEAR: 2026,
            PI_CALCINATION_METHOD: "A1",
            PI_CEMENT_PRODUCTION: 1_200_000,
            # Missing a1_kiln_feed_t_yr, a1_dust_return_correction_pct, a1_loi_raw_meal_pct
        }
        wb_bytes = _make_workbook(row)
        with pytest.raises(TemplateParseError) as exc_info:
            parse_excel(wb_bytes)
        cols = [e.column for e in exc_info.value.errors]
        assert "a1_kiln_feed_t_yr" in cols


# ---------------------------------------------------------------------------
# CSV parsing
# ---------------------------------------------------------------------------
class TestParseCsv:
    def _make_csv(self, row: dict) -> bytes:
        df = pd.DataFrame([row])
        buf = io.BytesIO()
        df.to_csv(buf, index=False)
        return buf.getvalue()

    def test_minimal_b1_csv(self):
        csv_bytes = self._make_csv(_base_plant_row())
        plant = parse_csv(csv_bytes)
        assert plant.plant_name == "Test Plant"
        assert plant.calcination_method == CalcinationMethod.B1_SIMPLE_OUTPUT

    def test_csv_no_fuels(self):
        csv_bytes = self._make_csv(_base_plant_row())
        plant = parse_csv(csv_bytes)
        assert plant.kiln_fuels == []
        assert plant.non_kiln_fuels == []

    def test_csv_missing_required(self):
        row = _base_plant_row()
        del row[PI_CEMENT_PRODUCTION]
        csv_bytes = self._make_csv(row)
        with pytest.raises(TemplateParseError):
            parse_csv(csv_bytes)


# ---------------------------------------------------------------------------
# parse_file dispatch
# ---------------------------------------------------------------------------
class TestParseFileDispatch:
    def test_xlsx_dispatched(self):
        wb_bytes = _make_workbook(_base_plant_row())
        plant = parse_file(wb_bytes, "data.xlsx")
        assert plant.plant_name == "Test Plant"

    def test_csv_dispatched(self):
        df = pd.DataFrame([_base_plant_row()])
        buf = io.BytesIO()
        df.to_csv(buf, index=False)
        plant = parse_file(buf.getvalue(), "data.csv")
        assert plant.plant_name == "Test Plant"

    def test_unsupported_extension(self):
        with pytest.raises(TemplateParseError) as exc_info:
            parse_file(b"data", "data.pdf")
        assert any("Unsupported" in e.message for e in exc_info.value.errors)


# ---------------------------------------------------------------------------
# End-to-end: parse → calculate
# ---------------------------------------------------------------------------
class TestEndToEnd:
    def test_full_pipeline(self):
        """Parse an Excel file and run the full carbon calculation."""
        from carbon_engine import CarbonEngine

        wb_bytes = _make_workbook(
            _base_plant_row(**{
                PI_ELEC_MWH: 50_000,
                PI_ELEC_EF: 0.5,
                PI_WATER_WITHDRAWAL: 1_200_000,
                PI_WATER_DISCHARGE: 200_000,
                PI_WATER_RECYCLED: 50_000,
            }),
            kiln_fuels=[
                _base_fuel_row(),
                _base_fuel_row(fuel_type="petrol_coke", consumption_t_per_yr=2_000, lhv_gj_per_t=32.5),
            ],
            non_kiln_fuels=[
                _base_fuel_row(fuel_type="diesel_oil", consumption_t_per_yr=500, lhv_gj_per_t=42.7),
            ],
            transport=[
                {TRANSPORT_DESC_COL: "Limestone", TRANSPORT_TKM_COL: 1_000_000, TRANSPORT_EF_COL: 0.062},
            ],
        )

        plant = parse_excel(wb_bytes)
        result = CarbonEngine(plant).calculate()

        assert result.total_co2e_t > 0
        assert result.scope1.calcination_co2_t > 0
        assert result.scope1.fuel_combustion_co2_t > 0
        assert result.scope2.electricity_co2_t == pytest.approx(25.0, rel=1e-3)
        assert result.scope3.transport_co2_t == pytest.approx(62.0, rel=1e-3)
        assert result.water.withdrawal_m3 == 1_200_000
        assert result.cbam.completeness_pct == 0.0
        assert len(result.source_mix) == 5
